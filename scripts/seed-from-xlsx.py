"""
One-time seed: turn SRR_Gutter_Pricing_Config.xlsx into the committed CSV
fallbacks in /data.

Run this only when the workbook itself changes shape. Day to day, the
fallbacks are refreshed from the live Google Sheet with `npm run snapshot`.

IMPORTANT: the `Unit Cost` column is deliberately dropped. /data is served
publicly by GitHub Pages, and Unit Cost plus target_margin would expose SRR's
margin. The estimator reads `Price` only.

Usage:  python scripts/seed-from-xlsx.py
"""
import csv
import os
import sys

try:
    import openpyxl
except ImportError:
    sys.exit("pip install openpyxl")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "SRR_Gutter_Pricing_Config.xlsx")
DATA = os.path.join(ROOT, "data")

PRICING_COLS = [
    "Code", "Name", "Description", "Material", "Category",
    "Profile", "UOM", "Price", "Active", "CostOfSaleAccount",
]

# Rule keys section 6 of the spec needs that are not yet in the workbook's
# Rules tab. Seeded here with the spec defaults so the geo math has something
# to read; they must also be added to the live sheet.
MISSING_RULES = [
    ("gutter_factor_gable", "0.62",
     "Share of perimeter that gets gutter on a gable roof (eaves only, 2 of 4 sides)."),
    ("gutter_factor_hip", "1.00",
     "Hip roof: all four sides are eaves, so the full perimeter gets gutter."),
    ("gutter_factor_flat", "0.95",
     "Flat or low slope roof."),
    ("gutter_factor_unknown", "0.80",
     "Used when the customer answers 'not sure' on roof type."),
    ("overhang_allowance_ft", "1.0",
     "Feet added per gutter-bearing side for the eave projecting past the wall line."),
    ("footprint_shape_factor", "1.12",
     "Method B only. Corrects 4*sqrt(area) for real houses not being perfect squares."),
]


def truthy(v):
    return "TRUE" if str(v).strip().upper() in ("TRUE", "1", "YES") else "FALSE"


def money(v):
    return "" if v is None else "{:.2f}".format(float(v))


# The Pricing tab stores Price and Active as formulas. openpyxl cannot evaluate
# formulas, and it drops Excel's cached results the moment anything re-saves
# the file, so reading `data_only` values alone is unreliable. This workbook
# only ever uses two formula shapes, so evaluate those two directly and refuse
# loudly on anything else rather than writing a blank price.
def resolve_price(value, formula, unit_cost, target_margin):
    if isinstance(value, (int, float)):
        return float(value)
    f = str(formula or "").replace(" ", "").upper()
    if f.startswith("=ROUND(") and "/(1-RULES!" in f:
        if unit_cost is None or target_margin is None:
            sys.exit("Cannot compute Price: missing Unit Cost or target_margin.")
        return round(float(unit_cost) / (1 - float(target_margin)), 2)
    sys.exit("Unrecognised Price formula, refusing to guess: {!r}".format(formula))


def resolve_active(value, formula):
    if isinstance(value, bool):
        return value
    if value is not None and str(value).strip() != "":
        return str(value).strip().upper() in ("TRUE", "1", "YES")
    f = str(formula or "").replace(" ", "").upper()
    if f in ("=TRUE()", "TRUE", "=TRUE"):
        return True
    if f in ("=FALSE()", "FALSE", "=FALSE"):
        return False
    sys.exit("Unrecognised Active formula, refusing to guess: {!r}".format(formula))


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    wbf = openpyxl.load_workbook(XLSX, data_only=False)

    # target_margin is needed only to recompute Price when Excel's cached
    # value is unavailable. It is never written out anywhere.
    target_margin = None
    for r in wbf["Rules"].iter_rows(min_row=2, values_only=True):
        if r[0] and str(r[0]).strip() == "target_margin":
            target_margin = float(r[1])

    # ---- Pricing -------------------------------------------------------
    ws = wb["Pricing"]
    wsf = wbf["Pricing"]
    rows = list(ws.iter_rows(values_only=True))
    frows = list(wsf.iter_rows(values_only=True))
    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    idx = {name: header.index(name) for name in header if name}

    for required in PRICING_COLS:
        if required not in idx:
            sys.exit("Pricing tab is missing required column: " + required)

    out = []
    for r, fr in zip(rows[1:], frows[1:]):
        code = r[idx["Code"]]
        # Stop at the first blank Code; everything past it is the legend text.
        if code is None or not str(code).strip():
            break
        rec = {}
        for col in PRICING_COLS:
            v = r[idx[col]]
            rec[col] = "" if v is None else str(v).strip()

        unit_cost = r[idx["Unit Cost"]] if "Unit Cost" in idx else None
        rec["Price"] = money(resolve_price(
            r[idx["Price"]], fr[idx["Price"]], unit_cost, target_margin))
        rec["Active"] = truthy(resolve_active(r[idx["Active"]], fr[idx["Active"]]))
        out.append(rec)

    with open(os.path.join(DATA, "pricing.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=PRICING_COLS, lineterminator="\n")
        w.writeheader()
        w.writerows(out)
    print("data/pricing.csv  {} rows ({} active)".format(
        len(out), sum(1 for r in out if r["Active"] == "TRUE")))

    # ---- Rules ---------------------------------------------------------
    ws = wb["Rules"]
    rrows = list(ws.iter_rows(values_only=True))
    rout = []
    seen = set()
    for r in rrows[1:]:
        key = r[0]
        if key is None or not str(key).strip():
            continue
        key = str(key).strip()
        seen.add(key)
        val = r[1]
        if isinstance(val, float) and val.is_integer():
            val = int(val)
        note = r[2] if len(r) > 2 and r[2] is not None else ""
        rout.append({"Key": key, "Value": str(val).strip(), "Notes": str(note).strip()})

    added = 0
    for key, val, note in MISSING_RULES:
        if key not in seen:
            rout.append({"Key": key, "Value": val,
                         "Notes": note + "  [NOT YET IN THE WORKBOOK - add it]"})
            added += 1

    # target_margin is internal. It is never read by the estimator and has no
    # business being on a public URL.
    rout = [r for r in rout if r["Key"] != "target_margin"]

    with open(os.path.join(DATA, "rules.csv"), "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["Key", "Value", "Notes"], lineterminator="\n")
        w.writeheader()
        w.writerows(rout)
    print("data/rules.csv    {} keys ({} seeded from the spec, not the workbook)".format(
        len(rout), added))


if __name__ == "__main__":
    main()
